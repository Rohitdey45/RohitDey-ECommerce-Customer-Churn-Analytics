"""Build report and README from executed notebook artifacts; run from repository root."""
from pathlib import Path
import json
import pandas as pd
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
ROOT=Path.cwd();OUT=ROOT/'outputs';s=json.loads((OUT/'project_summary.json').read_text())
TITLE='E-Commerce Customer Churn Prediction & Risk Analytics using RFM and Machine Learning'
d=Document();sec=d.sections[0];sec.top_margin=Inches(.75);sec.bottom_margin=Inches(.7);sec.left_margin=sec.right_margin=Inches(.8)
normal=d.styles['Normal'];normal.font.name='Calibri';normal.font.size=Pt(10.5);normal.paragraph_format.space_after=Pt(7)
for name in ['Title','Heading 1','Heading 2']:
 d.styles[name].font.name='Calibri';d.styles[name].font.color.rgb=RGBColor.from_string('123C55')
d.styles['Heading 1'].font.size=Pt(17);d.styles['Heading 2'].font.size=Pt(12)
header=sec.header.paragraphs[0];header.text='ROHIT DEY   /   CUSTOMER CHURN & RISK ANALYTICS';header.style='Caption'
footer=sec.footer.paragraphs[0];footer.alignment=2;footer.add_run('Academic project • Page ')
fld=OxmlElement('w:fldSimple');fld.set(qn('w:instr'),'PAGE');footer._p.append(fld)
def p(t,style=None):return d.add_paragraph(t,style)
def h(n,title):d.add_heading(f'{n:02d}  {title}',level=1)
def bullets(items):
 for t in items:p(t,'List Bullet')
def table(headers,rows):
 t=d.add_table(rows=1,cols=len(headers));t.style='Light Shading Accent 1'
 for c,v in zip(t.rows[0].cells,headers):c.text=str(v)
 for row in rows:
  for c,v in zip(t.add_row().cells,row):c.text=str(v)
 for row in t.rows:
  for cell in row.cells:
   for para in cell.paragraphs:
    for run in para.runs:run.font.size=Pt(9)
 # Repeat header if the table spans pages.
 trPr=t.rows[0]._tr.get_or_add_trPr();repeat=OxmlElement('w:tblHeader');trPr.append(repeat)
 return t
def chart(file,caption,width=6.3):
 d.add_picture(str(OUT/'charts'/file),width=Inches(width));p(caption,'Caption')
def page():d.add_page_break()
# 01 cover.
p('AICTE | IBM SkillsBuild', 'Subtitle');p('DATA ANALYTICS WITH AI\nACADEMIC INTERNSHIP 2026','Subtitle')
d.add_heading(TITLE,0)
p('FINAL PROJECT REPORT','Subtitle');p('Prepared by: Rohit Dey')
p('Conducted by BharatCares in association with AICTE and IBM SkillsBuild')
p('Submission edition • 23 September 2026')
p('Official practice data → audited reconstruction → customer analytics → leakage-safe experiment → decision support')
p('Source note: The supplied official practice dataset was available as a PDF export. With student approval, it was reconstructed and validated by page coordinates and text cross-checks. The original Excel workbook was not available; original sheet names and cell types are not claimed.')
p('Key outcome: A complete analytical workflow, but NOT a deployment-ready churn classifier. Held-out recall is zero at the specified 0.50 binary threshold. This limitation is a central result, not concealed by accuracy.')
page();d.add_heading('Report guide',0)
sections=['Cover Page','Student Details','Internship Details','Abstract','Introduction','Problem Statement','Objectives','Dataset Description','Dataset Source','Data Dictionary','Data Cleaning','Exploratory Data Analysis','Customer-Level Aggregation','RFM Analysis','Churn Definition','Target Leakage Prevention','Machine Learning Methodology','Logistic Regression','Model Evaluation','Confusion Matrix','Customer Churn Probability','Customer Risk Segmentation','Dashboard','Business Insights','Business Actions / Recommendations','Limitations','Future Scope','Conclusion','References']
for i,name in enumerate(sections,1):p(f'{i:02d}  {name}')
page();h(2,'Student Details');table(['Field','Detail'],[['Student name','Rohit Dey'],['Project title',TITLE],['Role','Student analyst / project author']])
p('No institution, student ID, mentor identity or completion claim has been invented. No private email address or class-chat participant information is included.')
h(3,'Internship Details');p('AICTE | IBM SkillsBuild Data Analytics with AI Academic Internship 2026, conducted by BharatCares in association with AICTE and IBM SkillsBuild. Supplied Masterclass 1–3 materials govern the learning workflow. The provided form screenshots require a notebook/code file, requirements.txt, Word report and Markdown README, plus an actual GitHub repository URL.')
h(4,'Abstract');p(f"This project reconstructs {s['raw_rows']:,} transaction records from the supplied official practice PDF, audits quality, and develops descriptive sales analysis and customer-level RFM features. After exact deduplication, {s['clean_master_rows']:,} records remain in the cleaned master. The sales view includes {s['total_orders']:,} orders with observed revenue of INR {s['total_revenue']:,.0f}; the separate identifiable customer-event view contains {s['total_customers']} customers. A snapshot at {s['cutoff']} separates historical features from a subsequent 180-day inactivity outcome. Of {s['cohort_customers']} eligible historical customers, {s['churned']} had no observed future purchase ({s['churn_rate']:.2%}). Logistic Regression is evaluated on {s['test_customers']} held-out customers. Accuracy is {s['metrics']['Accuracy']:.2%}, ROC-AUC is {s['metrics']['ROC_AUC']:.4f}, but precision, recall and F1 are zero at threshold 0.50. The dashboard therefore supports transparent exploration and further validation rather than automated retention decisions.")
h(5,'Introduction');p('E-commerce transaction records can describe past performance and inform customer-retention questions. RFM summarizes how recently, how often and how much customers purchase. Predictive usefulness requires an outcome observed after the information used to generate a prediction. The project follows Raw Data → Clean Data → EDA → Business Insights → Prediction → Dashboard → Business Decision.')
h(6,'Problem Statement');p('Identify customers known at a historical cutoff who may make no recorded purchase over the next 180 days, and assess whether Logistic Regression adds useful decision support beyond a majority-class prediction. Separately, describe product, category, regional and customer sales patterns without inferring unsupported causes.')
h(7,'Objectives');bullets(['Inspect the actual source structure and reconstruct without inventing records.','Document quality actions and preserve original data.','Create customer-level aggregates and interpretable RFM features.','Resolve inconsistent classroom churn definitions explicitly.','Train and evaluate a leakage-safe Logistic Regression pipeline.','Generate probabilities, unchanged classroom risk bands, and a dashboard.','Separate observations, hypotheses and evidence-limited business recommendations.'])
page();h(8,'Dataset Description')
table(['Measure / scope','Actual result'],[['Raw reconstructed shape','2,000 rows × 9 columns'],['Raw unique customers',s['raw_unique_customers']],['Raw unique orders',s['raw_unique_orders']],['Exact duplicates',s['duplicates']],['Valid date range',f"{s['date_start']} to {s['date_end']}"],['Cleaned master rows',s['clean_master_rows']],['Sales view rows / unique orders',s['sales_rows']],['Customer-event rows',s['event_rows']],['Full-period event customers',s['total_customers']],['Historical predictive cohort',s['cohort_customers']]])
p('The row grain is one transaction/order in this export after exact duplicate removal; remaining Order_ID values are unique. Region contains city labels rather than broad geographic territories. August ends on day 20 and is not a complete month. Currency markers in the source indicate INR. Original workbook sheet names and native types cannot be recovered from this PDF.')
h(9,'Dataset Source');p('Dataset: Official Masterclass 1 Practice Dataset provided with the IBM SkillsBuild Data Analytics with AI internship materials.')
p('Source supplied: '+s['source_filename']);p('PDF SHA-256: '+s['source_sha256'])
p('Pages 1–44 contain seven transaction columns; pages 45–88 contain Revenue and Profit; pages 89–95 contain classroom prompts. Every paired page has matching row baselines. The extractor decodes full PDF text cells, preserves missing cells, and compares the assembled cell text with independent standard text extraction. All 44 page-pair checks pass. Duplicate fingerprints also align across both halves. These checks establish internal PDF consistency, not certification against an unavailable original spreadsheet.')
p('A local derived workbook, official_practice_reconstructed.xlsx, contains a newly created Reconstructed_Transactions sheet. The notebook identifies that sheet by its actual headers and verifies a text-preserving Excel round trip. It is not represented as the original workbook.')
p('No verified public e-commerce dataset URL was found. The supermarket example’s link is not used; the chat’s generic Kaggle link is not a dataset source. README retains [INSERT OFFICIAL DATASET LINK IF PROVIDED BY TRAINER]. Raw data redistribution permission is unknown; source PDF and reconstructed data are excluded from the deliverable archive and Git tracking.')
h(10,'Data Dictionary')
dictionary=[['Order_ID','Identifier (string)','Order key; distinct order counts, not a model feature.'],['Order_Date','Date after parsing','Purchase date; strict DD-MM-YYYY; invalid values become NaT.'],['Customer_ID','Identifier (string)','Coded customer key; excluded from model inputs.'],['Product','Categorical','Product label; missing values explicitly Unknown.'],['Category','Categorical','Apparel, Beauty, Electronics, Furniture, Home after normalization.'],['Region','Categorical','Bangalore, Chennai, Delhi, Hyderabad, Kolkata, Mumbai, Pune; Unknown for missing.'],['Quantity','Numeric','Ordered units; negative values quarantined as ambiguous.'],['Revenue','Numeric, INR','Recorded transaction revenue; missing values not fabricated.'],['Profit','Numeric, INR','Recorded profit; negative profit retained, unknown profit omitted from sums.']]
table(['Actual source column','Analytical type','Meaning / use'],dictionary)
p('Imported reconstructed raw fields are strings. This preserves mixed number formatting and identifiers before deliberate conversion. Product values observed include Laptop, Phone, Headphones, Chair, Vacuum Cleaner, Camera, Bookshelf, Desk, Makeup Kit, Shoes, Jeans, Tablet, Speaker, Microwave, Skincare Set, Table, Sofa, Cookware Set, T-Shirt, Kurta, Jacket, Smartwatch, Blender and Perfume. Raw capitalization variants are printed in the notebook.')
page();h(11,'Data Cleaning')
table(['Column','Raw missing values'],s['missing_raw'].items())
audit=pd.read_csv(OUT/'cleaning_audit.csv');table(['Cleaning action','Affected','Reason'],audit.values.tolist())
p('Counts overlap; do not sum them as independent removals. Six missing products remain after duplicate removal, compared with seven raw missing products. The cleaned master retains all non-duplicate records; sales and customer views apply explicit inclusion rules. The sales view excludes invalid dates, missing revenue and suspect purchases, but includes valid sales with missing Customer_ID. The event view requires a known customer and valid date but permits missing revenue so purchase occurrence is not erased.')
p('Missing quantity does not automatically invalidate a known purchase; unknown quantity is excluded only from quantity sums. Negative quantities are not sign-flipped because they might represent returns. Fifteen negative-profit records are retained. Known mojibake currency text, INR markers and commas are removed before parsing; unrecognized formats stop execution. No median revenue is inserted into transaction records.')
out=pd.read_csv(OUT/'outlier_audit.csv');table(list(out.columns),out.round(2).values.tolist())
p('IQR outliers remain in the dataset; high-value electronics can legitimately be expensive. A source-row/page provenance file and quarantined sales file support local review. Twenty-three customers with any invalid-date record are excluded from the predictive cohort to reduce false inactivity labels; that exclusion can bias the sample.')
page();h(12,'Exploratory Data Analysis');bullets(s['observations'])
p(f"Sales-view observed revenue is INR {s['total_revenue']:,.0f} and average order value is INR {s['average_order_value']:,.2f}. These totals are not claimed to represent all original transactions: excluded/unknown amounts are not included. Profit sums omit {s['missing_profit_sales']} missing values; quantity charts omit {s['missing_quantity_sales']} missing quantities within the sales view.")
chart('revenue_by_month.png','Figure 1. Monthly revenue. August is partial; January-to-July endpoint comparison is not a monotonic trend claim.')
chart('revenue_by_category.png','Figure 2. Category contribution: observed revenue, not achievement against a business target.')
chart('revenue_by_region.png','Figure 3. Regional revenue. Unknown is a missing-data bucket, not an actual region.')
chart('top_products.png','Figure 4. Ten highest-revenue products in the sales view.')
chart('top_customers.png','Figure 5. Highest observed-revenue customers, shown using coded identifiers only.')
chart('quantity_by_category.png','Figure 6. Known units sold by category; missing quantities are omitted.')
chart('profit_by_category.png','Figure 7. Observed category profit; negative-profit transactions retained.')
chart('profit_by_region.png','Figure 8. Observed region profit; missing profit values are not zero-filled.')
chart('customer_distribution.png','Figure 9. One-time versus repeat customers in the full-period event view.')
chart('revenue_distribution.png','Figure 10. Transaction revenue distribution; unusually high values are not automatically deleted.')
h(13,'Customer-Level Aggregation');p('The groupby operation produces one row per Customer_ID. Order_Count is the number of distinct orders; Total_Revenue is the sum of available amounts; First_Purchase and Last_Purchase are the earliest/latest valid event dates. Revenue_Missing_Orders records incompleteness. Avg_Order_Value equals Total_Revenue / Order_Count only if every contributing order amount is present. Otherwise it remains missing. The notebook displays the first ten customers and column names.')
cs=pd.read_csv(OUT/'customer_rfm_full_period.csv');table(['Customer','Orders','Observed revenue (INR)','First purchase','Last purchase'],cs[['Customer_ID','Order_Count','Total_Revenue','First_Purchase','Last_Purchase']].head(10).values.tolist())
h(14,'RFM Analysis');table(['Feature','Business meaning','Implementation'],[['Recency','How long since the latest purchase','Reference date minus Last_Purchase, in days'],['Frequency','How often the customer bought','Order_Count, using unique Order_ID'],['Monetary','How much the customer generated','Total_Revenue if complete, otherwise missing']])
p('Full-period descriptive RFM references 20 August 2026. The model does not use these full-period features. Historical RFM is recomputed at 21 February 2026 using only earlier transactions. Missing historical Monetary and AOV are imputed inside the training pipeline; observed totals are never replaced in the source records.')
page();h(15,'Churn Definition');p('The Masterclass 3 workbook is internally inconsistent: page 12 gives Recency ≥90 days, then a combined rule of inactivity ≥180 days, Frequency <5 and Monetary <100,000. Page 15 separately recommends historical features and a future 180-day inactivity outcome. The definitions are not interchangeable.')
p('Selected definition: Churn_Status = 1 if an eligible customer with at least one purchase by the historical cutoff has no recorded valid purchase in the next 180 calendar days; 0 if the customer purchases at least once. This is an observed future inactivity proxy, not proof of permanent departure. No frequency or monetary criterion defines the label.')
table(['Period / rule','Implementation'],[['Feature period','1 January–21 February 2026 inclusive (52 calendar days)'],['Snapshot cutoff','21 February 2026'],['Outcome period','22 February–20 August 2026 inclusive (180 days)'],['Eligibility','Known purchase by cutoff; no invalid-date record for customer'],['Predictive cohort','204 customers'],['Inactive / returned','27 inactive; 177 returned'],['Observed inactivity rate',f"{s['churn_rate']:.2%}"],['Descriptive 90-day comparison','119 full-period customers; not used for model'],['Combined descriptive 180-day comparison','21 with known monetary value; not used for model']])
p('A 180-day future window fits the observed date span, so the workbook’s predictive setup can be implemented. The cost is only 52 days of feature history and a small positive class. The export is assumed complete across the window, but no source event-log audit is available. New customers first seen after the cutoff are not in the prediction cohort. No dates beyond the source’s valid coverage are invented.')
h(16,'Target Leakage Prevention');bullets(['Historical transactions are on/before cutoff; label transactions are strictly after cutoff.','Customer_ID, target labels, future order counts and full-period RFM are not predictors.','Recency is historical, not the future inactivity quantity used to label churn.','Train/test indices do not overlap, and one customer appears once in the snapshot.','Median imputation and scaling fit only training customers through Pipeline.','Assertions enforce time boundaries, feature allowlist and uniqueness.'])
p('The retrospective exclusion of invalid-date customers is a quality restriction, not an input feature, and can create selection bias. A customer-random test split at one snapshot is not a chronological backtest. A future deployment would require a separate time-based validation cohort.')
h(17,'Machine Learning Methodology');p('The allowlist is Recency, Frequency, Monetary and Avg_Order_Value. A stratified 75%/25% customer split with random_state=42 produces 153 training and 51 test customers. The test set has 44 returned and seven inactive customers. Missing historical money features receive training medians, followed by StandardScaler. No tuning is performed on the test labels. A majority-class DummyClassifier provides a comparison.')
h(18,'Logistic Regression');p('Logistic Regression models the association between scaled historical behaviour and the probability of future inactivity. The pipeline uses default L2 regularization (C=1), default class weights, random_state=42 and max_iter=2000. Predicted class uses probability ≥0.50. That binary threshold differs from the classroom risk bands. Correlation between Monetary and AOV limits independent coefficient interpretation. Coefficients do not establish causal effects.')
page();h(19,'Model Evaluation');metrics=pd.read_csv(OUT/'model_metrics.csv')
rows=[]
for _,r in metrics.iterrows():rows.append([r.Model]+[f'{r[c]:.2%}' for c in ['Accuracy','Precision','Recall','F1']]+[f'{r.ROC_AUC:.4f}' if pd.notna(r.ROC_AUC) else 'Not calculated'])
table(['Model','Accuracy','Precision','Recall','F1','ROC-AUC'],rows)
p('At threshold 0.50, the Logistic Regression model predicts no held-out customer as inactive. Its 86.27% accuracy equals the majority baseline. Precision is mathematically undefined with no positive predictions and is reported as zero using zero_division=0; recall and F1 are zero. ROC-AUC 0.6526 indicates some ranking signal in this small test set, but it does not establish calibration or operational benefit. No confidence interval or repeated temporal test was performed.')
p((OUT/'classification_report.txt').read_text())
chart('roc_curve.png','Figure 11. Held-out ROC curve, using probabilities; this does not rescue zero recall at threshold 0.50.')
h(20,'Confusion Matrix');table(['Actual / predicted','Returned','Inactive'],[['Returned',44,0],['Inactive',7,0]])
chart('confusion_matrix.png','Figure 12. Positive = future inactivity. TN=44, FP=0, FN=7, TP=0.')
p('A false positive contacts someone who would return anyway and may waste incentives. A false negative misses someone who becomes inactive: seven such cases occur here. Zero false positives is not evidence of a successful classifier when the model never predicts the positive class.')
h(21,'Customer Churn Probability');p('heldout_predictions.csv contains Customer_ID, Actual_Churn_Status, Predicted_Churn_Status and Churn_Probability for test customers. customer_predictions.csv scores the entire historical cohort and marks every row as Held-out or Training (in-sample). Only held-out predictions support reported metrics. The probabilities refer to the February historical snapshot, not a new live August forecast. Future outcomes for a live forecast are not available.')
h(22,'Customer Risk Segmentation');table(['Probability','Classroom band','Customers'],[['Below 40%','Low Risk',s['risk_counts']['Low Risk']],['40% to below 70%','Medium Risk',s['risk_counts']['Medium Risk']],['70% or higher','High Risk',s['risk_counts']['High Risk']]])
p('The top 20 are highest-ranked probabilities, not 20 High Risk customers. No customer reaches 70%. The two Medium Risk records are training observations, so their rankings are not independent validation evidence.')
top=pd.read_csv(OUT/'top20_risk_customers.csv');rows=[]
for _,r in top.iterrows():rows.append([r.Customer_ID,int(r.Recency),int(r.Frequency),f'{r.Monetary:,.0f}',f'{r.Churn_Probability:.2%}',r.Risk_Level,'Test' if r.Prediction_Set=='Held-out' else 'Train'])
table(['Customer','R days','F orders','M INR','Probability','Risk','Set'],rows)
d.add_heading('Observed top-20 patterns',2);bullets(s['risk_observations'])
d.add_heading('Possible insights, not causes',2);bullets(s['possible_insights'])
p('The actual results do not support a blanket claim of high Recency and low Monetary among top-ranked customers. Top-20 Recency is lower and Monetary higher than cohort medians. Nineteen have one purchase, but the cohort frequency median is also one. Avoid overstating low frequency as a differentiator without a fuller comparison.')
h(23,'Dashboard');p('dashboard/customer_churn_dashboard.html is a standalone interactive Plotly dashboard with embedded JavaScript, KPI cards, revenue trend, category/region performance, risk distribution, probability distribution, separate R/F/M distributions and a High Risk table. It requires no CDN or network access. Hover, zoom and legend interactions are available. The High Risk table explicitly states that no customer crosses the threshold.')
p('Sales/customer KPIs use full-period cleaned views; churn, churn rate and risk KPIs use the 204-customer historical cohort. These scopes are labelled directly in the interface. A static notebook KPI table and PNG charts are available if HTML rendering is restricted by the submission platform.')
h(24,'Business Insights');d.add_heading('Verified comparisons',2);bullets(s['insights']);d.add_heading('Hypotheses requiring additional evidence',2);bullets(s['hypotheses'])
h(25,'Business Actions / Recommendations');bullets(s['actions'])
p('Business decision: use this work as an analytical prototype and data-quality review, not automated churn intervention. Higher recall is useful when missing genuinely at-risk customers is costly; higher precision is useful when contact capacity or incentive budget is limited. With no supplied cost matrix or campaign capacity, a universal best threshold or ROI cannot be claimed. If capacity were constrained, validate ranking at that capacity on a separate validation set before selecting recipients. A randomized control group is required to estimate incremental campaign benefit.')
h(26,'Limitations');bullets(['Original Excel workbook unavailable: sheet names, formulas and native types are unverified; reconstruction is internally checked only.','Practice dataset provenance is official materials, but real-world representativeness and continuous observation are not established.','Invalid dates, missing amounts/IDs and ambiguous negative quantities cause view-specific exclusions and possible bias.','Only 52 pre-cutoff calendar days, 204 eligible customers and seven positive test outcomes.','Single random customer split at one cutoff; no temporal generalization or repeated-fold uncertainty estimate.','Recall is zero at threshold 0.50; accuracy only matches the baseline.','No full calibration validation; probabilities are not guaranteed event frequencies.','All-cohort risk outputs mix training and test scores with explicit labels.','August is incomplete; totals omit unknown revenue/profit and are not full audited financial accounts.','180-day inactivity can reflect seasonality or purchase cadence, not permanent customer loss.','No causal inference, motivations, retention uplift or ROI is established.'])
h(27,'Future Scope');bullets(['Obtain and reconcile the original workbook, missing dates and return semantics.','Collect longer purchase history and confirm observation completeness.','Backtest across several chronological cutoffs with disjoint future windows.','Evaluate calibrated probabilities, uncertainty and precision/recall at business capacity.','Choose alternate thresholds using validation data and documented error costs—not this test set.','Explore product-specific inactivity windows, then assess against the classroom definition transparently.','Run a small consent-based randomized retention pilot before expanding campaigns.'])
h(28,'Conclusion');p('The completed notebook, report and dashboard demonstrate the official cleaning, EDA, RFM, Logistic Regression, evaluation and risk workflow using executed results only. The key scientific conclusion is not that the model is ready to predict churn well: it misses all seven held-out inactive customers at the chosen threshold. Honest baseline comparison, leakage prevention and transparent limitations make the project reproducible and academically defensible. Better data and temporal evaluation are prerequisites to any live retention deployment.')
h(29,'References');bullets(['Official supplied AI_Powered_Data_Analytics_Theory_Deck - Masterclass 1.pdf, 14 pages: raw data, quality and analytical questions.','Official supplied AI_Powered_Data_Analytics_Theory_Deck - Masterclass 2.pdf, 16 pages: EDA, observations/insights and verification.','Official supplied Masterclass3_Student_Workbook.pdf, 31 pages: customer aggregation/RFM; p.12 conflicting churn rules; pp.14–18 leakage-safe modelling; pp.22–29 risk bands and actions.','Official supplied '+s['source_filename']+', 95 pages: source data and prompts.','Supermarket Sales Analysis DA project.pdf, 2 pages: structural example only; no numerical result or dataset copied.','GMT20260916-104325_RecordingnewChat.txt: participant chat context, not a verbatim lecture transcript. Generic Kaggle links are not dataset verification.','Supplied Google Form screenshots image-1.png, image-2.png, image-3.png: required files, upload sizes and GitHub repository field.','pandas documentation: https://pandas.pydata.org/docs/','scikit-learn documentation: https://scikit-learn.org/stable/','Plotly Python documentation: https://plotly.com/python/'])
p('AI assistance disclosure: AI-assisted code and documentation preparation was used. Numerical results originate from execution against the supplied practice PDF. Rohit Dey should review, understand and comply with institutional disclosure rules before submission. This document does not claim external certification or trainer approval.')
d.save(ROOT/'RohitDey_ProjectReport.docx')
# README uses the same executed summary.
structure='''RohitDey-ECommerce-Customer-Churn-Analytics/
├── RohitDey_ECommerceCustomerChurnPrediction.ipynb
├── requirements.txt
├── RohitDey_ProjectReport.docx
├── README.md
├── SUBMISSION_GUIDE.md
├── data/
│   └── README.md                 # place official PDF here locally; excluded from Git
├── dashboard/
│   └── customer_churn_dashboard.html
├── outputs/
│   ├── model_metrics.csv
│   ├── customer_risk_table.csv   # local/private; ignored by Git
│   ├── project_summary.json
│   ├── cleaning_audit.csv
│   ├── quality_checks.csv
│   └── charts/
├── scripts/
│   ├── create_documents.py
│   └── prepare_public_repository.py
└── .gitignore'''
readme=f'''# {TITLE}

**Author:** Rohit Dey  
**Internship:** AICTE | IBM SkillsBuild Data Analytics with AI Academic Internship 2026  
**Conducted by:** BharatCares in association with AICTE and IBM SkillsBuild

> **Result, not a success claim:** held-out accuracy is **{s['metrics']['Accuracy']:.2%}**, but recall, precision and F1 are **0** at threshold 0.50. The model predicts no positive test cases and matches the majority baseline. This is an academic prototype, **not a deployment-ready churn detector**.

## Overview and problem statement
Use the official practice transactions to understand sales and customer behaviour, then estimate which customers known at a historical cutoff will have **no recorded purchase during the next 180 days**. Use transparent evaluation before suggesting retention decisions.

## Objectives
- Preserve source data; audit every cleaning action.
- Explore monthly, product, category, region and customer patterns.
- Create one row per customer and RFM features.
- Resolve conflicting classroom labels and avoid future-information leakage.
- Fit Logistic Regression, evaluate against a baseline and calculate actual probabilities.
- Apply classroom risk bands; present an offline Plotly dashboard and cautious business actions.

## Dataset and dataset source
Dataset: Official Masterclass 1 Practice Dataset provided with the IBM SkillsBuild Data Analytics with AI internship materials.

**Verified public dataset link:** none found in supplied materials.  
**Trainer link placeholder:** [INSERT OFFICIAL DATASET LINK IF PROVIDED BY TRAINER]

The student could not upload Excel/ZIP and approved use of the supplied PDF:
`{s['source_filename']}`.

No Kaggle dataset is downloaded. The supermarket example is a structural reference only. Generic Kaggle links in chat do not establish a dataset source. The IBM Bob deck was not supplied and is not claimed as reviewed.

### PDF source exception
- 95 pages: transaction columns on pages 1–44, financial columns on pages 45–88, prompts on pages 89–95.
- Decode source PDF text objects and match printed row coordinates across all 44 page pairs.
- Cross-check headers, assembled text against standard extraction, and duplicate fingerprints; stop on mismatch.
- Preserve blank cells. No row value is fabricated.
- Create a **local reconstructed** workbook `data/official_practice_reconstructed.xlsx`, sheet `Reconstructed_Transactions`, and verify its Excel round trip.
- This does **not** verify original workbook sheets, formulas or native cell types. Raw imported dtypes are reconstructed strings, not original Excel dtypes.
- SHA-256: `{s['source_sha256']}`.

### Executed inspection
| Measure | Result |
|---|---:|
| Raw rows / columns | {s['raw_rows']:,} / 9 |
| Raw unique customers / orders | {s['raw_unique_customers']} / {s['raw_unique_orders']:,} |
| Exact duplicates | {s['duplicates']} |
| Valid date range | {s['date_start']} to {s['date_end']} |
| Cleaned master rows | {s['clean_master_rows']:,} |
| Sales-view orders | {s['total_orders']:,} |
| Full-period customer-event rows / customers | {s['event_rows']:,} / {s['total_customers']} |
| Sales-view observed revenue | INR {s['total_revenue']:,.0f} |
| Sales-view average order value | INR {s['average_order_value']:,.2f} |

August is incomplete. Missing money/excluded records are not included in these observed sales totals.

## Data dictionary
| Actual column | Analytical type | Meaning |
|---|---|---|
'''+''.join(f'| {a} | {b} | {c} |\n' for a,b,c in dictionary)+f'''
All actual source column names are preserved. Derived customer fields include `Order_Count`, `Total_Revenue`, `Avg_Order_Value`, `First_Purchase`, `Last_Purchase`, `Revenue_Missing_Orders`, `Recency`, `Frequency`, `Monetary`. Region labels are cities. Full raw categorical lists and dtypes are displayed in the notebook.

## Methodology
Raw Data → Clean Data → EDA → Business Insights → Prediction → Dashboard → Business Decision.

Customer-Level Dataset → RFM Features → Churn Target → Logistic Regression → Evaluation → Churn Probability → Risk Level → Business Action.

### Data cleaning
The original PDF and imported raw DataFrame remain unchanged. `outputs/cleaning_audit.csv` documents action counts and reasons. Exact duplicates alone are removed. Normalize case in Category/Region; use `Unknown` for missing Product/Region; parse dates strictly as DD-MM-YYYY; remove only recognized currency tokens and commas. Unknown numeric formats cause a fail-closed error.

24 invalid dates remain unavailable, never guessed. 41 ambiguous non-positive-quantity/negative-revenue records are quarantined from purchase analyses. 15 negative-profit rows remain valid loss observations. Outliers are flagged by IQR and retained. Missing transaction amounts are not imputed. Raw missing values: {s['missing_raw']}.

Sales EDA requires valid date/revenue and excludes suspect purchases; known customer IDs are not required. Customer events require valid date/ID and non-suspect purchases but may have missing revenue. Missing profit/quantity is omitted from the corresponding sum, not zero-filled. Audit counts overlap.

### EDA
Ten required sales/customer charts plus confusion matrix and ROC are generated with labelled axes. Questions include monthly growth, top products/customers, best/lowest-contributing categories and known regions, repeat purchases, quantity and observed profit. `Unknown` is not treated as a real region. Compare complete January–July endpoints; do not interpret partial August as a true decline.

### RFM and customer aggregation
One row per coded Customer_ID. Frequency counts unique orders. Recency uses the latest valid source date for descriptive RFM, and the historical cutoff for model RFM. Monetary is missing if any contributing revenue is absent. `Total_Revenue` is a known subtotal with an explicit completeness count. AOV is available only for complete amounts.

### Exact churn definition and inconsistency
Workbook p.12 contains both Recency ≥90 days and a combined 180-day/Frequency <5/Monetary <100,000 rule. Page 15 recommends historical features and a future 180-day inactivity target.

**Selected:** label 1 for an eligible historical customer with **no recorded valid purchase in (21 February 2026, 20 August 2026]**; otherwise 0. This is future inactivity, not proven permanent departure. Frequency/Monetary restrictions do not define the target.

- Historical features: **1 January–21 February 2026 inclusive (52 days)**.
- Outcome: **22 February–20 August 2026 inclusive (180 days)**.
- Cohort: **{s['cohort_customers']}** customers already observed by cutoff; customers with any invalid-date record excluded.
- Inactive: **{s['churned']} ({s['churn_rate']:.2%})**; returned: **{s['cohort_customers']-s['churned']}**.
- Descriptive alternatives, not used in training: 119 customers meet 90-day inactivity; 21 meet the known-value combined rule.

### Target leakage prevention
Historical features are calculated independently of future outcomes. Exclude Customer_ID, target, future counts and full-period RFM from model inputs. Split customers before fitting imputation/scaling. Assertions check temporal separation and disjoint customer splits. A single-snapshot random split is **not** temporal backtesting. Continuous source coverage is an unverified assumption.

### ML model and evaluation
Features: historical Recency, Frequency, Monetary, Avg_Order_Value. Pipeline: median SimpleImputer → StandardScaler → LogisticRegression (default L2, C=1, default class weights, max_iter=2000). Stratified 75%/25% split; random_state=42; {s['train_customers']} train and {s['test_customers']} test customers. Binary threshold 0.50, unchanged after evaluation.

| Held-out metric | Actual value |
|---|---:|
| Accuracy | {s['metrics']['Accuracy']:.2%} |
| Precision | 0.00% (zero_division=0; no positive predictions) |
| Recall | 0.00% |
| F1 | 0.0000 |
| ROC-AUC | {s['metrics']['ROC_AUC']:.4f} |
| Majority-class baseline accuracy | {s['baseline_accuracy']:.2%} |

Confusion matrix (rows actual returned/inactive, columns predicted returned/inactive):
```text
[[44, 0],
 [ 7, 0]]
```
All seven positive test cases are missed. ROC-AUC measures ranking only; calibration and campaign usefulness remain unproven.

## Predictions and risk levels
`heldout_predictions.csv` contains actual/predicted labels and probabilities for test customers. `customer_risk_table.csv` scores the historical cohort and labels each row **Held-out** or **Training (in-sample)**. These are February retrospective scores, not live August forecasts. Training scores are not evaluation evidence.

| Probability | Risk level | Count |
|---|---|---:|
| <40% | Low Risk | {s['risk_counts']['Low Risk']} |
| 40% to <70% | Medium Risk | {s['risk_counts']['Medium Risk']} |
| ≥70% | High Risk | {s['risk_counts']['High Risk']} |

“Top 20” means the largest model probabilities, not 20 High Risk customers. Scores range from 19.57% to 45.82%. Nineteen of these customers made one historical purchase; the cohort median is also one. Their median Recency is 13 vs 18 days, and Monetary is INR 148,327.50 vs 39,656. Do not impose the workbook example's high-recency/low-money pattern on these different results.

## Dashboard
Open `dashboard/customer_churn_dashboard.html` in a modern browser. It embeds Plotly JavaScript, needs no CDN, and provides hover/zoom/legend interactions. Seven KPI cards, revenue trend, category/region performance, risk/probability distributions, R/F/M distributions and a High Risk table are included. The latter is empty by the exact threshold and says so explicitly. KPI labels separate full-period sales/customer scope from the historical model cohort.

## Business insights
'''+''.join('- '+x+'\n' for x in s['observations']+s['insights'])+'''
### Possible hypotheses, not established explanations
'''+''.join('- '+x+'\n' for x in s['hypotheses'])+'''
### Business actions and decision
'''+''.join('- '+x+'\n' for x in s['actions'])+'''
False positives waste contact/incentive budget; false negatives miss inactive customers. Choose trade-offs using validation data and actual business costs. No cost ratio, guaranteed retention benefit or ROI is invented. Observed associations do not prove causation or customer motivations.

## Limitations and future scope
PDF-only reconstruction; short historical window; missing/invalid data and conservative selection; practice-data representativeness uncertain; only seven positive test cases; no temporal backtest; zero recall at 0.50; probabilities not fully calibrated; partial August; training rows in the all-cohort risk table. Obtain original data, longer history and multiple time cutoffs, then test calibration and cost-aware thresholds on validation data. Assess any retention uplift with a randomized pilot.

## Installation
Tested in Python 3.13. Use Python 3.11+ and the pinned dependencies for reproducibility.
```bash
python -m venv .venv
# Windows:
.venv\\Scripts\\activate
# macOS / Linux:
source .venv/bin/activate
python -m pip install -r requirements.txt
jupyter notebook
```

## Execution instructions
1. Obtain the **same official PDF** from your trainer/materials. No dataset is bundled or automatically downloaded.
2. Place it in `data/` with the exact supplied filename listed above, or beside the notebook.
3. Open `RohitDey_ECommerceCustomerChurnPrediction.ipynb` and select **Kernel → Restart & Run All**.
4. The notebook reconstructs and verifies the local workbook, executes analysis, and writes outputs/dashboard.
5. To regenerate report and README from executed outputs:
   ```bash
   python scripts/create_documents.py
   ```
6. Review `outputs/quality_checks.csv` and `outputs/project_summary.json`. Never overwrite results with teaching examples.

**Colab:** upload the notebook and PDF; run `%pip install pandas numpy matplotlib seaborn scikit-learn plotly openpyxl pypdf` in a temporary setup cell if needed, then run the notebook. The notebook searches its current directory as well as `data/`. Download generated artifacts before the Colab session ends. The document-generation script is optional for notebook execution and needs python-docx.

**Headless rerun from repository root:**
```bash
jupyter nbconvert --to notebook --execute --inplace RohitDey_ECommerceCustomerChurnPrediction.ipynb
```

The PDF adapter intentionally fails on a different layout. Do not feed an unrelated PDF or claim that its positional checks generalize to arbitrary spreadsheets. A newly supplied original workbook should be inspected independently rather than silently substituted.

## Project structure
```text
'''+structure+'''
```

## Privacy and GitHub publication
Raw PDF, raw/reconstructed Excel, cleaned transactions, row provenance and customer-level exports are not authorized for public redistribution by default. The student submission contains coded IDs and sample rows; it is not automatically a public-data license.

- Keep the full submission local/private; a private repository may require granting evaluator access.
- To create a **public-safe source-only copy**, run:
  ```bash
  python scripts/prepare_public_repository.py
  ```
  It clears notebook outputs and excludes the executed report, dashboard and all customer/data exports. It includes source code, setup/docs and aggregate metrics only. The repository README explains these intentional omissions. Do not upload the entire private submission ZIP to public GitHub without permission.
- This task prepares files; no GitHub repository has been created or uploaded. Enter the actual URL after creating it in your own account. No credentials are requested or embedded.
- Suggested repository description: “End-to-end e-commerce customer churn prediction and risk analytics project using RFM analysis, Logistic Regression, customer segmentation and interactive data visualization.”

## Submission and placeholders
Use `SUBMISSION_GUIDE.md` for exact Google Form upload mapping and checklist. The four mandatory files are the notebook, requirements.txt, RohitDey_ProjectReport.docx and README.md; HTML/ZIP are additional downloads, not replacements for these fields.

Remaining student action: create a repository, enter its real URL in the Google Form, and fill your personal form fields accurately. The dataset URL placeholder remains only until a verified trainer link is supplied; do not invent one. The report does not require an invented institution or mentor name.

## References and authorship
Primary references: supplied Masterclass 1 theory deck; Masterclass 2 theory deck; Masterclass 3 student workbook (especially pp.12–18 and 22–29); official practice PDF; Google Form screenshots. Supermarket reference is structural only. Chat is not a lecture transcript.

Software documentation: https://pandas.pydata.org/docs/ · https://scikit-learn.org/stable/ · https://plotly.com/python/

**Author: Rohit Dey.** AI assisted code and documentation preparation; computed results come from notebook execution. Review the work and follow institutional AI-disclosure rules before submission. No internship completion or trainer approval is claimed.
'''
(ROOT/'README.md').write_text(readme,encoding='utf-8')
print('Report and README generated from executed results.')
